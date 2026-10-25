#!/usr/bin/env python3
"""
S2 交付前静态检查。任何环境形态都先跑；它不替代编译、单测和环境验证，PASS 只表示没查出静态错误。

用法：
  python precheck.py --out <S2目录> --repo <仓或大仓子目录> [--repo ...]
                     [--handoff <改动包目录>] [--base <目录名>=<基线commit>] [--no-patch] [--allow-dirty]
                     [--write-manifest] [--config <precheck.json>] [--s1 <S1目录>]
  python precheck.py --lint-sop <sop-v2目录> [--config <precheck.json>]
  python precheck.py --ingest <原文文件> [<原文文件> ...] --out <S1/S1-00-原文>
  python precheck.py --selftest
  --out 放文档（上线手册、验证清单、S2-08-审查.md、S2-02-sql/、S2-03-config/）；--handoff 放 patches/README/manifest（离线快照 / 本机不可编译）。
  --s1 默认取 --out 的同级 S1/（读 S1-04-设计说明.md / S1-07-设计.md / S1-02-需求与验收.md）。

--repo 可以是独立 git 仓，也可以是大仓里的子目录：脚本用 git 定位仓根，只统计该目录下的改动。
报告列出每个目录检查了多少文件；同一大仓里有改动却没传进 --repo 的目录会被列出（U1）。
改动文件路径解析不到 → ERROR（P1）；没有任何改动文件 → ERROR（E0）。

检查项：
  H1 交付目录文件齐全              G1 无未提交文件（--allow-dirty 降 WARN）
  H2 交付目录不留 S2-04-verify/out.txt、S2-04-verify/env.local（自测输出与凭据）
  J1 Java import 可解析            J2 实体字段 ↔ S2-02-sql/ DDL 列（跳过 @TableField(exist=false)、static）
  F1 前端 import 可解析            F2 前端请求路径 ↔ Controller 映射
  Q1 上线手册含 S2-02-sql/*.sql 原文，或「文件名 + sha256:哈希」引用
  Q2 上线手册包含 S2-03-config/*.diff 的每条改动行
  Q3 SQL 可重入（CREATE TABLE 缺 IF NOT EXISTS、INSERT 无 NOT EXISTS/ON DUPLICATE/IGNORE/先删 → ERROR）
  N1 S2-03-config/*.diff 新增接口路径 ↔ Controller 映射
  N2 Feign 调用路径 ↔ 提供方 Controller 映射（含 ref_dirs）
  S1 Controller 入参 DTO 含身份字段（安全类，仅 WARN）
  C1 S1「改动面」列出的每个仓、以及带扩展名的文件都有改动       C2 代码走读「接触点」里要改/新建的仓都在改动面
  D1 新功能 S1 产物结构（评审摘要、接口调用者/数据范围/状态前置、按建议会怎样、手册术语映射）；精简版S1-07-设计.md 同样要有评审摘要、按建议会怎样、交付选项
  D2 决定表里状态为有效的 D-xx 必须出现在验收里；有设计说明就得有决定表
  R1 S2-08-审查.md：AC 全覆盖、结论用规定词、引用位置存在、环境登记的节名在上线手册里、阻塞类问题无「待修」
  V1 审查结论为已实现或环境登记的 AC 必须写进验证清单
  X1 S1/S1-00-原文/正文.md 为 UTF-8 且无替换字符；media/ 每张图在清单里有状态，待看和无法看不能过
  K1 类型为小改动、线上缺陷、纯配置时，不得另有新建表、平台登记或（纯配置）代码改动
  U1 大仓里未检查的改动目录        U2 改动包里的 patch 都对应一个 --repo
  T1 改动文件不留「待联调核对」类标记（config.defer_markers）
  A1 S2/S2-05-api/ 接口文档与冒烟集合的路径都在 Controller 映射里；集合覆盖 S2-04-verify/cases.json 的全部 AC
  M1 manifest 哈希（仅改动包形态）
  P1 改动文件路径不存在            E0 没有改动文件
  --lint-sop：L1 篇幅上限  L2 正文出现具体日期  L3 出现 project_terms

已知局限：外部依赖的类只能判断「依赖大概率存在」，类名对不对由编译确认；行为、口径、时效参数不在检查范围。
项目参数从 --config JSON 读，默认见 DEFAULT_CONFIG；放 artifacts/_project/S0-02-precheck.json，项目级，不按需求裁剪。
退出码 0 PASS / 1 FAIL。
"""
import argparse, hashlib, io, json, os, re, shutil, subprocess, sys, tempfile, zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

DEFAULT_CONFIG = {
    "api_prefixes": ["/api/"],
    "table_prefixes": [],
    "fe_aliases": {"@/": "src", "~/": "src"},
    "ignore_dirs": ["node_modules", "target", "dist", ".git", "build"],
    "ref_dirs": [],
    "identity_fields": ["mobile", "phone", "account", "userId", "userAccount"],
    "project_terms": [],
    "defer_markers": ["联调核对", "待联调", "需内网", "待核对", "待确认"],
}
# 与 sop-overview/SOP约束.md 一致。md 合计与 py 上限上调：原文落盘和决定表、验收覆盖、改动面文件、图清单写在本文件，精简版补评审摘要与交付选项，不另开文件。
LINT_LIMITS = {"SKILL.md": 120, "md": 200, "md_total": 1120, "py": 1280}
DATE_RE = re.compile(r'\b20\d{2}-(?:0[1-9]|1[0-2])-(?:0[1-9]|[12]\d|3[01])\b')
AC_RE = re.compile(r'\bAC-\d+\b')
VERDICTS = ("已实现", "未实现", "本期不做", "环境登记")
PLAN_ACTION_RE = re.compile(r'(?<!不)(改|新建)')
LOC_RE = re.compile(r'([\w./\\-]+\.(?:java|kt|ts|tsx|js|jsx|vue|xml|sql|py|go|cs)):(\d+)')
FEIGN_RE = re.compile(r'@FeignClient\s*\(([^)]*)\)')
BODY_TYPE_RE = re.compile(r'@RequestBody\s+(?:@\w+(?:\([^)]*\))?\s+)*(?:final\s+)?([\w.]+)')
JDK_PREFIX = ("java.", "javax.", "jakarta.")
FE_EXT = (".ts", ".tsx", ".js", ".jsx", ".vue")
FE_RESOLVE = ["", ".ts", ".tsx", ".js", ".jsx", ".vue", "/index.ts", "/index.tsx", "/index.js", "/index.vue"]
MAPPING_ANN_RE = re.compile(r'@(Request|Get|Post|Put|Delete|Patch)Mapping\s*\(([^)]*)\)')
STR_RE = re.compile(r'"([^"]+)"')
IMPORT_JAVA_RE = re.compile(r'^\s*import\s+(?:static\s+)?([\w.]+?)(\.\*)?\s*;', re.M)
IMPORT_FE_RE = re.compile(r'(?:import\s+[^;\'"]*?from\s*|import\s*\(?\s*|require\s*\(\s*)[\'"]([^\'"]+)[\'"]')
FE_CONST_RE = re.compile(r'(?:const|let|var)\s+(\w+)\s*=\s*[\'"`]([^\'"`\n]+)[\'"`]')
TABLE_RE = re.compile(r'@TableName\s*\(\s*(?:value\s*=\s*)?"([^"]+)"')
FIELD_RE = re.compile(r'((?:@\w+(?:\([^)]*\))?\s*)*)(?:private|protected)\s+((?:static\s+|final\s+)*)[\w<>\[\], ?]+\s+(\w+)\s*(?:=[^;]*)?;')
DDL_RE = re.compile(r'CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?`?(\w+)`?\s*\((.*?)\)\s*ENGINE', re.S | re.I)
DDL_COL_RE = re.compile(r'^\s*`(\w+)`', re.M)


class Report:
    def __init__(self):
        self.issues = []
        self.counts = {}  # 目录名 -> {"java": n, "fe": n, "other": n}

    def err(self, c, m): self.issues.append(("ERROR", c, m))
    def warn(self, c, m): self.issues.append(("WARN", c, m))
    def info(self, c, m): self.issues.append(("INFO", c, m))
    def count(self, repo, kind): self.counts.setdefault(repo, {"java": 0, "fe": 0, "other": 0})[kind] += 1

    @property
    def changed(self): return sum(sum(v.values()) for v in self.counts.values())

    def finish(self):
        n_err = sum(1 for i in self.issues if i[0] == "ERROR")
        n_warn = sum(1 for i in self.issues if i[0] == "WARN")
        lines = [f"{lv} [{c}] {m}" for level in ("ERROR", "WARN", "INFO") for lv, c, m in self.issues if lv == level]
        per = "；".join(f"{r} java {v['java']} / 前端 {v['fe']} / 其他 {v['other']}" for r, v in self.counts.items()) or "无"
        lines.append(f"检查文件：{per}")
        lines.append(f"改动文件 {self.changed}；ERROR {n_err}；WARN {n_warn}")
        lines.append("RESULT: " + ("PASS" if n_err == 0 else "FAIL"))
        return "\n".join(lines), n_err == 0


def read(p):
    try: return Path(p).read_text(encoding="utf-8", errors="ignore")
    except Exception: return ""


def git(cwd, *args):
    r = subprocess.run(["git", "-C", str(cwd), *args], capture_output=True, text=True, encoding="utf-8", errors="ignore")
    return r.stdout.rstrip() if r.returncode == 0 else ""  # 不能 lstrip：status --porcelain 首行以空格开头


def locate(repo_dir):
    """返回 (git根, 前缀)。repo_dir 可以是仓根或子目录。"""
    top = git(repo_dir, "rev-parse", "--show-toplevel")
    if not top: return None, None
    prefix = git(repo_dir, "rev-parse", "--show-prefix")
    return Path(top), prefix


def changed_all(root, base):
    """返回相对 git 根的改动文件集合与未提交文件列表。"""
    out, dirty = set(), []
    if base:
        out.update(git(root, "diff", "--name-only", "--diff-filter=ACMR", f"{base}..HEAD").splitlines())
    for line in git(root, "status", "--porcelain", "-uall").splitlines():
        if len(line) > 3:
            path = line[3:].split(" -> ")[-1].strip().strip('"')
            dirty.append(path)
            if "D" not in line[:2]: out.add(path)
    return {f for f in out if f}, dirty


def java_imports(txt):
    """[(名字, 是否通配)]；通配时名字是包名。"""
    return [(m.group(1), bool(m.group(2))) for m in IMPORT_JAVA_RE.finditer(txt)]


def camel_to_snake(s): return re.sub(r'(?<!^)(?=[A-Z])', '_', s).lower()


def norm_sql(s):
    return re.sub(r'[ \t]+', ' ', "\n".join(l.rstrip() for l in s.strip().splitlines()))


def index_repo(repo_dir, cfg, idx):
    """收集 FQCN、import、Controller 映射、已有表名。"""
    for root, dirs, files in os.walk(repo_dir):
        dirs[:] = [d for d in dirs if d not in cfg["ignore_dirs"]]
        for fn in files:
            p = Path(root) / fn
            if fn.endswith(".java"):
                txt = read(p)
                m = re.search(r'^\s*package\s+([\w.]+)\s*;', txt, re.M)
                if m: idx["fqcn"][f"{m.group(1)}.{fn[:-5]}"] = p
                idx["simple"].setdefault(fn[:-5], set()).add(p.resolve())
                for name, wild in java_imports(txt):
                    idx["imports"].setdefault(name + (".*" if wild else ""), set()).add(p)
                if FEIGN_RE.search(txt): continue  # Feign 接口是调用方，不算提供方映射
                cls_prefix = ""
                head = txt.split("class ", 1)[0]
                cm = MAPPING_ANN_RE.search(head)
                if cm and cm.group(1) == "Request":
                    s = STR_RE.findall(cm.group(2))
                    cls_prefix = "/" + s[0].strip("/") if s else ""
                    if not s and cm.group(2).strip(): idx["const_mappings"] += 1
                for kind, args in MAPPING_ANN_RE.findall(txt):
                    strs = STR_RE.findall(args)
                    if not strs and args.strip(): idx["const_mappings"] += 1
                    for s in strs:
                        if kind == "Request" and cls_prefix.strip("/") == s.strip("/"): continue
                        idx["mappings"].add(cls_prefix + "/" + s.strip("/"))
                idx["tables"].update(TABLE_RE.findall(txt))
            elif fn.endswith(".xml") and cfg["table_prefixes"]:
                pat = r'\b((?:' + "|".join(map(re.escape, cfg["table_prefixes"])) + r')\w+)\b'
                idx["tables"].update(re.findall(pat, read(p)))


def imported_elsewhere(idx, exclude, pred):
    return any(pred(k) and any(p not in exclude for p in ps) for k, ps in idx["imports"].items())


def find_src_root(p, alias_dir):
    for parent in p.parents:
        if parent.name == alias_dir: return parent
    return None


def check_java_import(rel, name, wild, idx, src_packages, exclude, rep):
    if name.startswith(JDK_PREFIX): return
    if wild:
        if name in src_packages: return
        if imported_elsewhere(idx, exclude, lambda k: k.startswith(name + ".")): return
    else:
        if name in idx["fqcn"]: return
        parent = name.rsplit(".", 1)[0]
        if parent in idx["fqcn"]: return  # 内部类 / 静态导入
        if imported_elsewhere(idx, exclude, lambda k: k == name): return
        if parent in src_packages:
            rep.err("J1", f"{rel}: import {name} 的包在工作区有源码，但该类不存在且无他处 import（拼错或臆造）"); return
    stem = ".".join(name.split(".")[:3])
    if imported_elsewhere(idx, exclude, lambda k: k.startswith(stem + ".")): return  # 同组依赖已在用
    rep.warn("J1", f"{rel}: import {name}{'.*' if wild else ''} 无源文件，同组依赖也未见使用（新依赖？确认 pom）")


def md_table_after(text, heading_kw, header=False):
    """标题含 heading_kw 的小节里第一张表的数据行（header=True 时含表头）；无该小节返回 None。"""
    lines = text.splitlines()
    start = next((i for i, l in enumerate(lines) if l.lstrip().startswith("#") and heading_kw in l), None)
    if start is None: return None
    rows = []
    for l in lines[start + 1:]:
        s = l.strip()
        if s.startswith("#"): break
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if not all(re.fullmatch(r':?-{2,}:?', c) or not c for c in cells): rows.append(cells)
        elif rows and s: break
    return rows if header else rows[1:]


def check_sql_reentrant(sf, rep):
    txt = re.sub(r'--[^\n]*', '', read(sf))
    deleted, seen = set(), set()
    for s in (x.strip() for x in re.split(r';\s*(?:\n|$)', txt)):
        if not s: continue
        u = s.upper()
        m = re.match(r'DELETE\s+FROM\s+`?(\w+)`?', s, re.I)
        if m: deleted.add(m.group(1).lower()); continue
        if re.match(r'CREATE\s+TABLE\b', s, re.I) and not re.match(r'CREATE\s+TABLE\s+IF\s+NOT\s+EXISTS', s, re.I):
            rep.err("Q3", f"{sf.name}: CREATE TABLE 缺 IF NOT EXISTS，重跑报错")
        m = re.match(r'INSERT\s+(IGNORE\s+)?INTO\s+`?(\w+)`?', s, re.I)
        if m and not (m.group(1) or "NOT EXISTS" in u or "ON DUPLICATE KEY" in u or m.group(2).lower() in deleted):
            if m.group(2) not in seen:
                seen.add(m.group(2))
                rep.err("Q3", f"{sf.name}: INSERT INTO {m.group(2)} 重跑会重复或主键冲突（加 NOT EXISTS / ON DUPLICATE KEY，或先删后插）")
        if re.match(r'ALTER\s+TABLE\b', s, re.I) and re.search(r'\bADD\b', s, re.I) and "INFORMATION_SCHEMA" not in txt.upper():
            rep.warn("Q3", f"{sf.name}: ALTER TABLE ... ADD 重跑会报错（先查 information_schema）")


def check_feign(rel, txt, mappings, rep):
    fm = FEIGN_RE.search(txt)
    if not fm: return
    head, _, body = txt.partition("interface ")
    pm = re.search(r'path\s*=\s*"([^"]*)"', fm.group(1))
    cm = MAPPING_ANN_RE.search(head)
    base = pm.group(1) if pm else (STR_RE.findall(cm.group(2)) or [""])[0] if cm else ""
    base = "/" + base.strip("/") if base.strip("/") else ""
    for _, args in MAPPING_ANN_RE.findall(body):
        for s in STR_RE.findall(args):
            p = re.sub(r'\{[^}]*\}', '{}', base + "/" + s.strip("/"))
            if not any(m == p or m.endswith(p) for m in mappings):
                rep.warn("N2", f"{rel}: Feign 路径 {p} 在提供方 Controller 映射中未找到（路径写错，或提供方不在 --repo / ref_dirs）")


def check_identity(rel, txt, idx, cfg, rep):
    if not re.search(r'@(Rest)?Controller\b', txt): return
    for t in sorted(set(BODY_TYPE_RE.findall(txt))):
        for p in sorted(idx["simple"].get(t.rsplit(".", 1)[-1], ())):
            fields = {f for _, mods, f in FIELD_RE.findall(read(p)) if "static" not in mods}
            hit = sorted(fields & set(cfg["identity_fields"]))
            if hit: rep.warn("S1", f"{rel}: 入参 {t} 含身份字段 {','.join(hit)}，身份应取自登录态（安全类，记入审查「后续优化」）")


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
FILE_TOKEN = re.compile(r'[\w./\\-]+\.(?:java|kt|ts|tsx|js|jsx|vue|xml|sql|py|go|cs|yml|yaml|json)')
NARROW_TYPE = ("小改动", "纯配置", "线上缺陷")


def _all_tables(text):
    tables, cur = [], []
    for line in text.splitlines():
        s = line.strip()
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if all(re.fullmatch(r':?-{2,}:?', c) or not c for c in cells): continue
            cur.append(cells)
        elif cur:
            tables.append(cur); cur = []
    if cur: tables.append(cur)
    return tables


def requirement_type(s1):
    if not s1.is_dir(): return ""
    blobs = [read(p) for p in (s1.parent / "00-当前位置.md", s1 / "S1-07-设计.md", s1 / "S1-04-设计说明.md") if p.exists()]
    m = re.search(r'类型[:：]\s*([^\s（(]+)', "\n".join(blobs))
    return m.group(1) if m else ""


def narrow_type(kind):
    return any(kind.startswith(k) for k in NARROW_TYPE)


def check_type(s1, sql_files, changed, rep):
    kind = requirement_type(s1)
    if not narrow_type(kind): return
    if kind.startswith("纯配置") and changed:
        rep.err("K1", "类型是纯配置，但检测到代码改动。有代码改动不能标纯配置")
    if any(re.search(r'CREATE\s+TABLE\b', read(f), re.I) for f in sql_files):
        rep.err("K1", f"类型是{kind}，但 S2-02-sql/ 里有新建表。新建表应按新功能做全套，不能靠类型跳过")
    walk = s1 / "S1-03-代码走读.md"
    if not walk.exists(): return
    rows = md_table_after(read(walk), "平台登记", header=True)
    if not rows: return
    col = next((i for i, h in enumerate(rows[0]) if "要加" in h or "本需求" in h), None)
    if col is None: return
    for cells in rows[1:]:
        v = cells[col].strip() if col < len(cells) else ""
        if v and v not in ("无", "—", "-", "不适用"):
            rep.err("K1", f"类型是{kind}，平台登记要新增内容。有平台登记应按新功能做全套，不能靠类型跳过")
            return


def check_decisions(s1, rep):
    dec = s1 / "S1-01-决定表.md"
    if not dec.exists():
        if (s1 / "S1-04-设计说明.md").exists():
            rep.err("D2", "有 S1-04-设计说明.md 但没有 S1-01-决定表.md，有效原句没有登记")
        return
    table = next((t for t in _all_tables(read(dec)) if t and "编号" in t[0] and "状态" in t[0]), None)
    if not table:
        rep.err("D2", "S1-01-决定表.md 没有「编号 / 状态」表"); return
    id_i, st_i = table[0].index("编号"), table[0].index("状态")
    acc = "\n".join(read(s1 / n) for n in ("S1-02-需求与验收.md", "S1-07-设计.md", "S1-04-设计说明.md") if (s1 / n).exists())
    for cells in table[1:]:
        did = cells[id_i].strip() if id_i < len(cells) else ""
        status = cells[st_i].strip() if st_i < len(cells) else ""
        if not re.fullmatch(r'D-\d+', did): continue
        if "作废" in status or "替代" in status or not status.startswith("有效"): continue
        if did not in acc:
            rep.err("D2", f"决定表 {did} 状态为有效，验收里没有引用它")


def _section_empty(txt, name):
    lines = txt.splitlines()
    for i, line in enumerate(lines):
        if name in line and "〔图" in line:
            a = next((j for j in range(i, -1, -1) if lines[j].lstrip().startswith("#")), 0)
            b = next((j for j in range(i + 1, len(lines)) if lines[j].lstrip().startswith("#")), len(lines))
            body = "\n".join(lines[a + 1:b])
            body = re.sub(r'〔图[^〕]*〕', '', body)
            return not body.strip()
    return False


def check_source(s1, rep):
    if not s1.is_dir(): return
    root = s1 / "S1-00-原文"
    body, manifest = root / "正文.md", root / "清单.md"
    if not body.exists():
        rep.err("X1", "S1/S1-00-原文/正文.md 不存在。原文先经 precheck.py --ingest 落盘，决定表只引用该文件"); return
    try:
        txt = body.read_bytes().decode("utf-8")
    except UnicodeDecodeError:
        rep.err("X1", "S1/S1-00-原文/正文.md 不是 UTF-8"); return
    if txt.startswith("\ufeff") or "\ufffd" in txt:
        rep.err("X1", "S1/S1-00-原文/正文.md 含 BOM 或替换字符")
    if not manifest.exists():
        rep.err("X1", "S1/S1-00-原文/清单.md 不存在"); return
    rows = md_table_after(read(manifest), "图片", header=True) or []
    listed = {}
    if len(rows) > 1:
        hi = {h: i for i, h in enumerate(rows[0])}
        fi, si, pi = hi.get("文件"), hi.get("状态"), hi.get("位置")
        for cells in rows[1:]:
            if fi is None or fi >= len(cells): continue
            listed[cells[fi].strip()] = (
                cells[si].strip() if si is not None and si < len(cells) else "",
                cells[pi].strip() if pi is not None and pi < len(cells) else "")
    media = root / "media"
    if media.is_dir():
        for f in sorted(media.iterdir()):
            if f.is_file() and f"media/{f.name}" not in listed and f.name not in listed:
                rep.err("X1", f"media/{f.name} 在清单里没有状态")
    for name, (status, place) in listed.items():
        if status == "待看":
            rep.err("X1", f"{name} 仍是待看。读图后改为已读；空节旁的图不能忽略，忽略了要进待定决策")
        elif status == "无法看":
            rep.err("X1", f"{name} 无法看，不能当作已读原文")
        elif status == "忽略" and place not in ("页眉", "页脚") and _section_empty(txt, name):
            rep.err("X1", f"{name} 旁边的节是空的，不能忽略。读图或进待定决策")
        elif status not in ("已读", "忽略"):
            rep.err("X1", f"{name} 的状态无效")


def _decode_text(path):
    raw = path.read_bytes()
    try:
        return raw.decode("utf-8-sig"), "utf-8"
    except UnicodeDecodeError:
        return raw.decode("gbk"), "gbk"


def _ptext(node):
    return "".join((t.text or "") for t in node.iter(W_NS + "t"))


def _hlevel(p):
    style = p.find(f"{W_NS}pPr/{W_NS}pStyle")
    val = (style.get(W_NS + "val") if style is not None else "") or ""
    if "Heading" in val or "标题" in val:
        m = re.search(r"(\d)", val)
        return int(m.group(1)) if m else 1
    ol = p.find(f"{W_NS}pPr/{W_NS}outlineLvl")
    if ol is not None and (ol.get(W_NS + "val") or "").isdigit():
        return int(ol.get(W_NS + "val")) + 1
    return 0


def _table_md(tbl):
    rows = []
    for tr in tbl.findall(W_NS + "tr"):
        cells = [_ptext(tc).replace("|", "/").replace("\n", " ") for tc in tr.findall(W_NS + "tc")]
        if any(cells): rows.append("| " + " | ".join(cells) + " |")
    if not rows: return ""
    ncol = max(r.count("|") - 1 for r in rows)
    return "\n".join([rows[0], "| " + " | ".join("---" for _ in range(ncol)) + " |", *rows[1:]])


def _rels(z, name):
    if name not in z.namelist(): return {}
    root = ET.fromstring(z.read(name))
    return {rel.get("Id"): rel.get("Target") for rel in root}


def _save_media(z, target, dest_dir, images, place):
    target = target.lstrip("/")
    path = target if target.startswith("word/") else "word/" + target.lstrip("./")
    if path not in z.namelist(): return ""
    ext = Path(path).suffix.lower() or ".bin"
    n = len(images) + 1
    name = f"b{n}{ext or '.bin'}"
    blob = z.read(path)
    (dest_dir / name).write_bytes(blob)
    status, reason = ("待看", "") if ext in (".png", ".jpg", ".jpeg", ".gif", ".webp") else ("无法看", "矢量图或未知格式，需另存为 png")
    if place in ("页眉", "页脚"): status, reason = "忽略", place
    images.append((f"media/{name}", place, status, reason))
    return f"〔图 media/{name}〕"


def _emit_part(z, xml_name, rel_name, dest_dir, images, place, out_lines):
    if xml_name not in z.namelist(): return
    rels = _rels(z, rel_name)
    root = ET.fromstring(z.read(xml_name))
    body = root.find(W_NS + "body")
    nodes = list(body) if body is not None else list(root.iter(W_NS + "p"))
    for node in nodes:
        if node.tag == W_NS + "tbl":
            md = _table_md(node)
            if md: out_lines.append(md)
            continue
        if node.tag != W_NS + "p": continue
        level, text, marks = _hlevel(node), _ptext(node).strip(), []
        for blip in node.iter(A_NS + "blip"):
            rid = blip.get(R_NS + "embed")
            if rid and rels.get(rid):
                marks.append(_save_media(z, rels[rid], dest_dir, images, place))
        if level and text: out_lines.append("#" * min(level, 6) + " " + text)
        elif text: out_lines.append(text)
        out_lines.extend(m for m in marks if m)


def _embed_xlsx(z, out_lines, missed):
    try:
        import openpyxl
    except ImportError:
        missed.append(("内嵌表格", "未安装 openpyxl")); return
    for name in z.namelist():
        if not (name.startswith("word/embeddings/") and name.endswith(".xlsx")): continue
        wb = openpyxl.load_workbook(io.BytesIO(z.read(name)), data_only=True, read_only=True)
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                vals = ["" if c is None else str(c).replace("|", "/") for c in row]
                if any(vals): rows.append("| " + " | ".join(vals) + " |")
            if not rows: continue
            ncol = max(r.count("|") - 1 for r in rows)
            out_lines.append(f"### {ws.title}")
            out_lines.append("\n".join([rows[0], "| " + " | ".join("---" for _ in range(ncol)) + " |", *rows[1:]]))
        wb.close()


def _ingest_docx(path, dest, images, missed):
    lines = []
    media = dest / "media"
    media.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path) as z:
        _emit_part(z, "word/document.xml", "word/_rels/document.xml.rels", media, images, "正文", lines)
        rels = _rels(z, "word/_rels/document.xml.rels")
        extra = []
        for target in rels.values():
            base = Path(target).name
            if base.startswith("header"):
                _emit_part(z, "word/" + target.lstrip("./"), "word/_rels/" + base + ".rels", media, images, "页眉", extra)
            elif base.startswith("footer"):
                _emit_part(z, "word/" + target.lstrip("./"), "word/_rels/" + base + ".rels", media, images, "页脚", extra)
        if extra:
            lines += ["", "## 页眉页脚", *extra]
        if "word/comments.xml" in z.namelist():
            comments = []
            _emit_part(z, "word/comments.xml", "word/_rels/comments.xml.rels", media, images, "批注", comments)
            if any(x.strip() for x in comments):
                lines += ["", "## 批注", *comments]
        before = len(lines)
        lines += ["", "## 内嵌表格"]
        _embed_xlsx(z, lines, missed)
        if lines[-1] == "## 内嵌表格" and len(lines) == before + 2:
            lines.pop(); lines.pop()
    return lines


def _ingest_doc(path, dest, images, missed):
    docx = dest / "_converted.docx"
    word = None
    try:
        import win32com.client
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        d = word.Documents.Open(str(Path(path).resolve()))
        d.SaveAs(str(docx.resolve()), FileFormat=16)
        d.Close(False)
    except Exception:
        missed.append((path.name, "doc 转 docx 失败，请另存为 docx"))
        return []
    finally:
        if word is not None:
            try: word.Quit()
            except Exception: pass
    try:
        return _ingest_docx(docx, dest, images, missed)
    finally:
        if docx.exists(): docx.unlink()


def _ingest_pdf(path, missed):
    try:
        from pdfminer.high_level import extract_text
        from pdfminer.pdfpage import PDFPage
    except ImportError:
        missed.append((path.name, "未安装 pdfminer")); return None
    try:
        with open(path, "rb") as fh:
            npages = sum(1 for _ in PDFPage.get_pages(fh))
    except Exception:
        missed.append((path.name, "pdf 无法读取")); return None
    lines, any_text = [], False
    for page in range(npages):
        text = (extract_text(str(path), page_numbers=[page]) or "").strip()
        if text:
            any_text = True
            lines += [f"## 第{page + 1}页", text]
        else:
            missed.append((f"第{page + 1}页", "无文本层"))
    if not any_text:
        missed.append((path.name, "整份无文本层")); return None
    return lines


def _write_manifest(dest, images, missed, ok, notes):
    rows = ["# 清单", "", "| 项 | 值 |", "|----|----|", f"| 结果 | {'PASS' if ok else 'FAIL'} |", *notes, "",
            "## 图片", "", "| 文件 | 位置 | 状态 | 原因 |", "|------|------|------|------|"]
    rows += [f"| {a} | {b} | {c} | {d} |" for a, b, c, d in images]
    rows += ["", "## 未提取", "", "| 对象 | 原因 |", "|------|------|"]
    rows += [f"| {a} | {b} |" for a, b in missed]
    (dest / "清单.md").write_text("\n".join(rows) + "\n", encoding="utf-8", newline="\n")


def ingest(paths, out_dir):
    """把原文落成 UTF-8 的 正文.md / media/ / 清单.md。标准输出只有 ASCII 的 RESULT 行。"""
    dest = Path(out_dir)
    dest.mkdir(parents=True, exist_ok=True)
    (dest / "media").mkdir(exist_ok=True)
    images, missed, body, notes = [], [], [], []
    for raw in paths:
        path = Path(raw)
        ext = path.suffix.lower()
        body += ["", f"## 来源 {path.name}", ""]
        if ext in (".md", ".txt"):
            text, enc = _decode_text(path)
            notes.append(f"| 编码 {path.name} | {enc} |")
            body.append(text.strip())
        elif ext == ".docx":
            body.extend(_ingest_docx(path, dest, images, missed))
        elif ext == ".doc":
            got = _ingest_doc(path, dest, images, missed)
            if got: body.extend(got)
        elif ext == ".pdf":
            got = _ingest_pdf(path, missed)
            if got: body.extend(got)
            elif not any(path.name in a for a, _ in missed):
                missed.append((path.name, "整份无文本层"))
        else:
            missed.append((path.name, "不支持的格式"))
    fatal = ("转 docx 失败", "不支持的格式", "无法读取", "整份无文本层", "未安装")
    content = [x for x in body if x and not str(x).startswith("## 来源")]
    ok = any(str(x).strip() for x in content) and not any(k in b for _, b in missed for k in fatal)
    if ok:
        text = "\n\n".join(x for x in body if x is not None).strip() + "\n"
        (dest / "正文.md").write_text(text, encoding="utf-8", newline="\n")
    elif (dest / "正文.md").exists():
        (dest / "正文.md").unlink()
    _write_manifest(dest, images, missed, ok, notes)
    print("RESULT: PASS" if ok else "RESULT: FAIL")
    return 0 if ok else 1


def check_s1_shape(s1, doc, rep):
    txt = read(doc)
    if not re.search(r'^#+.*评审摘要', txt, re.M): rep.err("D1", "S1-04-设计说明.md 缺「评审摘要」（S1 评审先发它）")
    if not any(l.lstrip().startswith("|") and all(k in l for k in ("调用者", "数据范围", "状态前置")) for l in txt.splitlines()):
        rep.err("D1", "S1-04-设计说明.md 接口契约缺 调用者/数据范围/状态前置 列")
    pend = s1 / "S1-05-待定决策.md"
    if pend.exists():
        ptxt = read(pend)
        if "按建议会怎样" not in ptxt: rep.err("D1", "S1-05-待定决策.md 缺「按建议会怎样」列")
        _delivery_options(ptxt, "S1-05-待定决策.md", rep)
    flow = s1 / "S1-06-交互流程.md"
    api = md_table_after(txt, "接口契约", header=True)
    pcol = next((i for i, h in enumerate(api[0]) if "路径" in h), None) if api else None
    if flow.exists() and pcol is not None:
        arrows = "\n".join(l for l in read(flow).splitlines() if ">>" in l)
        for cells in api[1:]:
            p = re.split(r'[?\s\[{]', cells[pcol].strip("` ") if pcol < len(cells) else "")[0]
            if p and p not in arrows: rep.err("D1", f"S1-06-交互流程.md 没有把接口 {p} 画进时序图箭头（覆盖表登记不算）")
    hb = s1.parent.parent / "_project" / "S0-01-项目手册.md"
    if hb.exists() and not re.search(r'^#+.*术语映射', read(hb), re.M): rep.err("D1", "S0-01-项目手册.md 缺「术语映射」小节")


def _delivery_options(text, where, rep):
    opts = md_table_after(text, "交付选项", header=True)
    col = next((i for i, h in enumerate(opts[0]) if "用户选择" in h), None) if opts else None
    if col is None:
        rep.err("D1", f"{where} 缺「交付选项」表或其「用户选择」列（选项要用户回复，不能由执行者代填结论）"); return
    for cells in opts[1:]:
        v = cells[col].strip() if col < len(cells) else ""
        if not v or re.search(r'^[（(]?待|待\S{0,8}(回复|确认|定)', v):
            rep.err("D1", f"交付选项「{cells[0]}」没有用户选择（待回复不算；依赖项和交付范围同一次问，S1 出口前问完）")
    for k in ("环境形态", "交付范围"):
        if not any(k in c[0] for c in opts[1:]): rep.err("D1", f"交付选项缺「{k}」行")


def check_slim_shape(doc, rep):
    """小改动、线上缺陷、纯配置的S1-07-设计.md：不套用新功能的接口三列，但这三项不能靠类型跳过。"""
    txt = read(doc)
    if not re.search(r'^#+.*评审摘要', txt, re.M): rep.err("D1", "S1-07-设计.md 缺「评审摘要」（S1 评审先发它）")
    if "按建议会怎样" not in txt: rep.err("D1", "S1-07-设计.md 缺「按建议会怎样」列")
    _delivery_options(txt, "S1-07-设计.md", rep)


def check_plan(s1, changed_by_repo, repos, rep):
    doc = next((s1 / f for f in ("S1-04-设计说明.md", "S1-07-设计.md") if (s1 / f).exists()), None)
    if doc is None: rep.warn("C1", f"{s1} 下没有 S1-04-设计说明.md / S1-07-设计.md，跳过改动面核对"); return
    if doc.name == "S1-04-设计说明.md": check_s1_shape(s1, doc, rep)
    else: check_slim_shape(doc, rep)
    rows = md_table_after(read(doc), "改动面")
    if not rows: rep.err("C1", f"{doc.name} 没有「改动面」表（它是 S2 的任务清单）"); return
    walk = md_table_after(read(s1 / "S1-03-代码走读.md"), "接触点", header=True) if (s1 / "S1-03-代码走读.md").exists() else None
    if walk:
        planned = " ".join(" ".join(c) for c in rows)
        col = next((i for i, h in enumerate(walk[0]) if "做法" in h), None)
        names = {p.name for d, root, _ in repos for b in {d.parent, root.parent} if b.is_dir() for p in b.iterdir() if p.is_dir()}
        seen = set()
        for cells in walk[1:]:
            if col is None or col >= len(cells) or not PLAN_ACTION_RE.search(cells[col]): continue
            for tok in sorted(set(re.findall(r'[A-Za-z][\w.]*(?:-[\w.]+)+', " ".join(cells))) - seen):
                hit = next((n for n in names if n == tok or n.startswith(tok + "-")), None)
                if hit and tok not in planned and hit not in planned:
                    seen.add(tok)
                    rep.err("C2", f"代码走读接触点「{cells[0]}」要改 {tok}，改动面没有（补进改动面，或写「{tok} | 不改：原因」）")
    for cells in rows:
        repo = cells[0].strip("` /").replace("\\", "/")
        # 只有「不改：原因」才跳过；叙述里出现「无改动」不能当豁免（防「precheck 按无改动登记」绕过）
        if not repo or any(re.search(r'不改\s*[:：]', c) or c.strip() == "不改" for c in cells): continue
        if not any(files and (repo in d.as_posix() or any(("/" + f).startswith("/" + repo + "/") or f"/{repo}/" in "/" + f for f in files))
                   for d, files in changed_by_repo):
            rep.err("C1", f"改动面列出的 {repo} 没有改动（漏改，或未传 --repo；无 git 先 init；确实不改写「不改：原因」）")
            continue
        changed = next((files for d, files in changed_by_repo if repo in d.as_posix()
                        or any(("/" + f).startswith("/" + repo + "/") or f"/{repo}/" in "/" + f for f in files)), [])
        for tok in FILE_TOKEN.findall(" ".join(cells)):
            base = Path(tok.replace("\\", "/")).name
            if not any(Path(f).name == base for f in changed):
                rep.err("C1", f"改动面列出的文件 {base} 在 {repo} 的改动里没有")


def check_api(out, mappings, cfg, rep):
    api = out / "S2-05-api"
    if not api.is_dir(): return
    def known(p):
        norm = re.sub(r'\{\{[^}]*\}\}|\{[^}]*\}|:\w+', '{}', p)
        pos = [i for i in (norm.find(pf) for pf in cfg["api_prefixes"]) if i >= 0]
        tail = norm[min(pos):] if pos else norm
        return any(m == tail or m.endswith(tail) or tail.endswith(m) for m in mappings)
    def items(node):
        for it in node.get("item", []):
            if "request" in it: yield it
            yield from items(it)
    coll_txt = ""
    for f in sorted(api.glob("*.json")):
        try: j = json.loads(read(f))
        except ValueError as e: rep.err("A1", f"api/{f.name} 不是合法 JSON：{e}"); continue
        for p in j.get("paths", {}) if isinstance(j, dict) else {}:
            if not known(p): rep.err("A1", f"api/{f.name}: 接口 {p} 在 Controller 映射中未找到（文档与代码不一致）")
        if f.name.endswith(".postman_collection.json"):
            coll_txt += read(f)
            for it in items(j):
                u = it["request"].get("url", "") if isinstance(it["request"], dict) else ""
                raw = u.get("raw", "") if isinstance(u, dict) else str(u)
                path = re.sub(r'^(\{\{[^}]+\}\}|https?://[^/]+)', '', raw.strip()).split("?")[0]
                if any(pf in path for pf in cfg["api_prefixes"]) and not known(path):
                    rep.err("A1", f"api/{f.name}: 用例「{it.get('name', '')}」请求 {path} 在 Controller 映射中未找到")
    cases = out / "S2-04-verify" / "cases.json"
    if coll_txt and cases.exists():
        miss = sorted(set(AC_RE.findall(read(cases))) - set(AC_RE.findall(coll_txt)), key=lambda x: int(x[3:]))
        if miss: rep.err("A1", f"冒烟集合未覆盖 verify 里的 {', '.join(miss)}")


def walk_files(dirs, cfg):
    out = []
    for d in dirs:
        for r, ds, fs in os.walk(d):
            ds[:] = [x for x in ds if x not in cfg["ignore_dirs"]]
            out += [Path(r) / f for f in fs]
    return out


def check_review(s1, out, repo_dirs, cfg, rep):
    src = next((s1 / f for f in ("S1-02-需求与验收.md", "S1-07-设计.md") if (s1 / f).exists()), None)
    if src is None: rep.warn("R1", f"{s1} 下没有 S1-02-需求与验收.md / S1-07-设计.md，跳过审查核对"); return
    review = out / "S2-08-审查.md"
    if not review.exists(): rep.err("R1", "缺 S2/S2-08-审查.md（独立审查未做）"); return
    rtxt = read(review)
    table = md_table_after(rtxt, "AC 实现映射", header=True) or [[]]
    col = next((i for i, h in enumerate(table[0]) if "结论" in h), -1)
    cj = out / "S2-04-verify" / "cases.json"
    verified = {a for v in re.findall(r'"ac"\s*:\s*"([^"]*)"', read(cj) if cj.exists() else "") for a in AC_RE.findall(v)}
    mapped = {}
    for cells in table[1:]:
        m = AC_RE.search(cells[0])
        if m: mapped[m.group(0)] = cells
    for ac in sorted(set(AC_RE.findall(read(src))), key=lambda x: int(x[3:])):
        cells = mapped.get(ac)
        if cells is None: rep.err("R1", f"S2-08-审查.md「AC 实现映射」缺 {ac}"); continue
        row, verdict = " ".join(cells), cells[col] if -len(cells) <= col < len(cells) else ""
        if not verdict.startswith(VERDICTS): rep.err("R1", f"{ac} 结论「{verdict[:20]}」不是 {'/'.join(VERDICTS)} 之一")
        elif verdict.startswith("未实现"): rep.err("R1", f"{ac} 未实现")
        elif verdict.startswith("已实现") and not LOC_RE.search(row): rep.err("R1", f"{ac} 标已实现但没有 文件:行")
        elif verdict.startswith("环境登记"):
            titles = re.findall(r'[「《]([^」》]+)[」》]', row)
            if "上线手册" not in row:
                rep.err("R1", f"{ac} 标环境登记但没写上线手册里的登记步骤")
            elif not titles:
                rep.err("R1", f"{ac} 环境登记要写上线手册里的节名（用「」标出）")
            else:
                heads = re.findall(r'^#+\s*(.+?)\s*$', read(out / "S2-06-上线手册.md") if (out / "S2-06-上线手册.md").exists() else "", re.M)
                for t in titles:
                    if not any(t.strip() == h or t.strip() in h for h in heads):
                        rep.err("R1", f"{ac} 环境登记指向的「{t}」在上线手册里没有这一节")
        elif verdict.startswith("本期不做") and ac in verified:
            rep.err("R1", f"{ac} 标本期不做但 verify 在测它：本端部分按已实现/未实现判，只有整条由另一端表现才本期不做")
    checklist = read(out / "S2-07-验证清单.md") if (out / "S2-07-验证清单.md").exists() else ""
    for ac, cells in mapped.items():
        verdict = cells[col] if -len(cells) <= col < len(cells) else ""
        if verdict.startswith(("已实现", "环境登记")) and ac not in checklist:
            rep.err("V1", f"{ac} 审查结论是已实现或环境登记，S2-07-验证清单.md 没有这条。清单全过盖不住没列入的验收")
    files = walk_files(repo_dirs, cfg)
    for path, line in sorted(set(LOC_RE.findall(rtxt))):
        tail = "/" + path.replace("\\", "/").lstrip("./")
        hit = next((f for f in files if f.as_posix().endswith(tail)), None)
        if hit is None or len(read(hit).splitlines()) < int(line):
            rep.err("R1", f"S2-08-审查.md 引用的位置 {path}:{line} 不存在")
    for cells in md_table_after(rtxt, "问题") or []:
        row = " ".join(cells)
        if re.search(r'需求不符|交付与部署', row) and "待修" in row:
            rep.err("R1", f"审查问题未关闭（需求不符、交付与部署须本轮修完）：{' '.join(cells[:3])}")


def lint_sop(sop_dir, cfg):
    root, issues, total = Path(sop_dir), [], 0
    terms = [t for t in cfg.get("project_terms", []) if t]
    for p in sorted(root.rglob("*")):
        if not p.is_file() or p.suffix not in (".md", ".py") or any(x in cfg["ignore_dirs"] for x in p.parts): continue
        rel, txt = p.relative_to(root).as_posix(), read(p)
        n = len(txt.splitlines())
        limit = LINT_LIMITS["py"] if p.suffix == ".py" else LINT_LIMITS["SKILL.md"] if p.name == "SKILL.md" else LINT_LIMITS["md"]
        if n > limit: issues.append(("L1", f"{rel} {n} 行，超过上限 {limit}"))
        if p.suffix == ".md" and p.name != "README.md": total += n
        m = DATE_RE.search(txt)
        if m: issues.append(("L2", f"{rel}: 出现具体日期 {m.group(0)}（修订历史写进变更说明，不进正文）"))
        low = txt.lower()
        for t in terms:
            if t.lower() in low: issues.append(("L3", f"{rel}: 出现项目专有内容「{t}」"))
    if total > LINT_LIMITS["md_total"]: issues.append(("L1", f".md 合计 {total} 行（不含 README），超过上限 {LINT_LIMITS['md_total']}"))
    lines = [f"ERROR [{c}] {m}" for c, m in issues]
    lines += [f"SOP 规则文本 {total} 行（上限 {LINT_LIMITS['md_total']}）；project_terms {len(terms)} 项；ERROR {len(issues)}",
              "RESULT: " + ("PASS" if not issues else "FAIL")]
    return "\n".join(lines), not issues


def load_cfg(path, rep=None):
    cfg = dict(DEFAULT_CONFIG)
    if path:
        try: cfg.update(json.loads(read(path)))
        except Exception as e:
            if rep: rep.err("CFG", f"读取 {path} 失败：{e}")
    return cfg


def run(a):
    rep = Report()
    cfg = load_cfg(a.config, rep)
    out = Path(a.out)
    handoff = Path(a.handoff) if getattr(a, "handoff", None) else (None if a.no_patch else out)
    bases = dict(b.split("=", 1) for b in a.base)

    for f in ["S2-06-上线手册.md", "S2-07-验证清单.md"]:
        if not (out / f).exists(): rep.err("H1", f"交付目录缺 {f}")
    if handoff is not None:
        if not (handoff / "README.md").exists(): rep.err("H1", "改动包缺 README.md")
        if not ((handoff / "patches").exists() and list((handoff / "patches").glob("*.patch"))):
            rep.err("H1", "改动包缺 patches/*.patch")
    for f in ("S2-04-verify/out.txt", "S2-04-verify/env.local", "S2-05-api/env.postman.json"):
        if (out / f).exists(): rep.err("H2", f"交付目录里还有 {f}（自测输出或凭据不进交付物，移到交付目录外）")
    sql_files = sorted((out / "S2-02-sql").glob("*.sql")) if (out / "S2-02-sql").exists() else []
    sql_files = [f for f in sql_files if f.name != "reset_test.sql"]

    repos = []
    for r in a.repo:
        d = Path(r).resolve()
        root, prefix = locate(d)
        if root is None: rep.err("G0", f"{d} 不在 git 仓内"); continue
        repos.append((d, root, prefix))
    if handoff is not None and (handoff / "patches").exists():
        for pf in sorted((handoff / "patches").glob("*.patch")):
            if not any(pf.stem == Path(r).resolve().name for r in a.repo):
                rep.err("U2", f"改动包里的 {pf.name} 没有对应的 --repo，这个仓的改动没有被检查")

    idx = {"fqcn": {}, "simple": {}, "imports": {}, "mappings": set(), "tables": set(), "const_mappings": 0}
    for d, _, _ in repos: index_repo(d, cfg, idx)
    for r in cfg["ref_dirs"]:
        if Path(r).is_dir(): index_repo(Path(r), cfg, idx)
        else: rep.warn("CFG", f"ref_dirs 里的 {r} 不存在")
    norm_mappings = {re.sub(r'\{[^}]*\}', '{}', m) for m in idx["mappings"]}
    src_packages = {k.rsplit(".", 1)[0] for k in idx["fqcn"]}
    if idx["const_mappings"]: rep.warn("F2", f"{idx['const_mappings']} 处 Controller 映射使用常量，未解析，F2/N1 可能漏报")

    ddl = {}
    for sf in sql_files:
        for t, body in DDL_RE.findall(read(sf)): ddl[t] = set(DDL_COL_RE.findall(body))

    api_pat = re.compile(r'[\'"`]([^\'"`\n]*(?:' + "|".join(map(re.escape, cfg["api_prefixes"])) + r')[A-Za-z0-9_\-/{}:.$]*)[\'"`]')
    fe_files = fe_paths_found = 0
    manifest_rows = []
    root_changes = {}  # git根 -> (全部改动, 已覆盖前缀)

    changed_by_repo = []
    for d, root, prefix in repos:
        allc, dirty_all = changed_all(root, bases.get(d.name))
        root_changes.setdefault(root, (allc, set()))[1].add(prefix)
        files = sorted(f[len(prefix):] for f in allc if f.startswith(prefix))
        dirty = [x for x in dirty_all if x.startswith(prefix)]
        if dirty: (rep.warn if a.allow_dirty else rep.err)("G1", f"{d.name} 有 {len(dirty)} 个未提交文件（交付前必须 commit）")
        changed_by_repo.append((d, files))

    all_changed_paths = {d / f for d, files in changed_by_repo for f in files}
    for d, files in changed_by_repo:
        rep.counts.setdefault(d.name, {"java": 0, "fe": 0, "other": 0})
        for rel in files:
            p = d / rel
            if not p.exists(): rep.err("P1", f"{d.name}/{rel} 不存在（路径解析失败）"); continue
            manifest_rows.append(f"{d.name}  {rel}  {hashlib.sha256(p.read_bytes()).hexdigest()}")
            txt = read(p)
            mk = next((m for m in cfg["defer_markers"] if m and m in txt), None)
            if mk: rep.err("T1", f"{d.name}/{rel} 留有「{mk}」：查代码证据定下来，定不了进待定决策，不留到联调")
            if rel.endswith(".java"):
                rep.count(d.name, "java")
                for name, wild in java_imports(txt):
                    check_java_import(rel, name, wild, idx, src_packages, all_changed_paths, rep)
                for t in TABLE_RE.findall(txt):
                    if t in ddl:
                        for ann, mods, field in FIELD_RE.findall(txt):
                            if "static" in mods or re.search(r'exist\s*=\s*false', ann): continue
                            cm = re.search(r'@Table(?:Field|Id)\s*\(\s*(?:value\s*=\s*)?"([^"]+)"', ann)
                            col = cm.group(1) if cm else camel_to_snake(field)
                            if col not in ddl[t]: rep.err("J2", f"{rel}: 字段 {field}→列 {col} 不在 S2-02-sql/ DDL `{t}`")
                    elif t not in idx["tables"]: rep.warn("J2", f"{rel}: @TableName {t} 既不在 S2-02-sql/ DDL 也不在仓内已有表中")
                check_feign(rel, txt, norm_mappings, rep)
                check_identity(rel, txt, idx, cfg, rep)
            elif rel.endswith(FE_EXT):
                rep.count(d.name, "fe"); fe_files += 1
                for imp in IMPORT_FE_RE.findall(txt):
                    base = None
                    if imp.startswith("."): base = p.parent / imp
                    else:
                        for alias, adir in cfg["fe_aliases"].items():
                            if imp.startswith(alias):
                                sr = find_src_root(p, adir)
                                if sr: base = sr / imp[len(alias):]
                    if base is not None and not any(Path(str(base) + e).exists() for e in FE_RESOLVE):
                        rep.err("F1", f"{rel}: import '{imp}' 解析不到文件")
                consts = dict(FE_CONST_RE.findall(txt))
                resolved = re.sub(r'\$\{(\w+)\}', lambda m: consts.get(m.group(1), m.group(0)), txt)
                found = set(api_pat.findall(resolved))
                found = {x for x in found if not any(q != x and q.startswith(x + "/") for q in found)}  # 去掉仅作前缀的常量
                for path in found:
                    fe_paths_found += 1
                    norm = re.sub(r'\$\{[^}]*\}|:\w+', '{}', path)
                    pos = min(i for i in (norm.find(pf) for pf in cfg["api_prefixes"]) if i >= 0)
                    tail = norm[pos:]
                    if not any(m == tail or m.endswith(tail) or tail.endswith(m) for m in norm_mappings):
                        rep.warn("F2", f"{rel}: 请求路径 {path} 在 Controller 映射中未找到")
            else:
                rep.count(d.name, "other")
    if rep.changed == 0 and not requirement_type(Path(a.s1) if getattr(a, "s1", None) else out.parent / "S1").startswith("纯配置"):
        rep.err("E0", "没有检测到任何改动文件，不允许交付空包（检查 --repo / --base）")
    if fe_files and fe_paths_found == 0: rep.warn("F2", "前端改动文件里未识别到任何请求路径，检查 config.api_prefixes")

    for root, (allc, prefixes) in root_changes.items():
        if "" in prefixes: continue
        missed = sorted({f.split("/", 1)[0] for f in allc if "/" in f and not any(f.startswith(px) for px in prefixes)})
        if missed: rep.warn("U1", f"{root.name} 里这些目录有改动但未传 --repo，没有被检查：{', '.join(missed)}")

    manual = read(out / "S2-06-上线手册.md")
    blocks = [norm_sql(b) for b in re.findall(r'```sql\s*\n(.*?)```', manual, re.S)]
    for sf in sql_files:
        check_sql_reentrant(sf, rep)
        raw = read(sf); c = norm_sql(raw)
        if any(c == b or c in b for b in blocks): continue
        h = hashlib.sha256(sf.read_bytes()).hexdigest()
        if sf.name in manual and f"sha256:{h}" in manual: continue
        rep.err("Q1", f"上线手册里没有 {sf.name} 的原文，也没有「{sf.name} sha256:{h}」引用")

    s1 = Path(a.s1) if getattr(a, "s1", None) else out.parent / "S1"
    if s1.is_dir():
        check_plan(s1, changed_by_repo, repos, rep)
        check_review(s1, out, [d for d, _, _ in repos], cfg, rep)
        check_decisions(s1, rep)
        check_source(s1, rep)
        check_type(s1, sql_files, rep.changed, rep)
    else: rep.warn("C1", f"未找到 S1 目录 {s1}，跳过改动面与审查核对（C1/R1）")
    check_api(out, norm_mappings, cfg, rep)

    cdir = out / "S2-03-config"
    if cdir.exists():
        manual_norm = re.sub(r'\s+', ' ', manual)
        pat = re.compile(r'((?:' + "|".join(map(re.escape, cfg["api_prefixes"])) + r')[A-Za-z0-9_\-/{}.*]+)')
        for df in sorted(cdir.glob("*")):
            added = [l[1:] for l in read(df).splitlines() if l.startswith("+") and not l.startswith("+++")]
            missing = [l.strip() for l in added if l.strip() and re.sub(r'\s+', ' ', l.strip()) not in manual_norm]
            if missing: rep.err("Q2", f"上线手册缺 {df.name} 的 {len(missing)} 条改动行，例：{missing[0]}")
            for line in added:
                for path in pat.findall(line):
                    if path.rstrip("/") not in idx["mappings"] and not any(m.startswith(path.rstrip("/").rstrip("*").rstrip("/")) for m in idx["mappings"]):
                        rep.warn("N1", f"{df.name}: 新增路径 {path} 在 Controller 映射中未找到")

    mf = (handoff or out) / "manifest.txt"
    if a.write_manifest:
        mf.parent.mkdir(parents=True, exist_ok=True)
        mf.write_text("\n".join(manifest_rows) + "\n", encoding="utf-8"); rep.info("M1", f"已生成 manifest.txt，{len(manifest_rows)} 个文件")
    elif handoff is not None:
        if not mf.exists(): rep.err("M1", "缺 manifest.txt（加 --write-manifest 生成）")
        else:
            have = set(read(mf).splitlines())
            for row in manifest_rows:
                if row not in have: rep.err("M1", f"manifest 与工作区不一致：{row.split()[1]}")
    return rep


def selftest():
    T = Path(tempfile.mkdtemp(prefix="precheck-selftest-"))
    try:
        mono = T / "mono"; svc = mono / "svc-a"; web = mono / "web-a"; other = mono / "svc-b"
        (svc / "src/main/java/com/x/ctl").mkdir(parents=True); (svc / "src/main/java/com/x/util").mkdir(parents=True)
        (web / "src/api").mkdir(parents=True); (web / "src/libs").mkdir(parents=True); other.mkdir()
        (svc / "src/main/java/com/x/util/Helper.java").write_text("package com.x.util;\npublic class Helper {}\n", encoding="utf-8")
        (svc / "src/main/java/com/x/ctl/OldController.java").write_text(
            'package com.x.ctl;\nimport com.x.util.Helper;\nimport org.springframework.web.bind.annotation.*;\nimport cn.tool.core.util.StrUtil;\n'
            '@RequestMapping("svcapi/old")\npublic class OldController { @PostMapping("/_search") public void s(){} }\n', encoding="utf-8")
        (web / "src/libs/axios.ts").write_text("export const defHttp={}", encoding="utf-8")
        (other / "a.txt").write_text("x", encoding="utf-8")
        g = lambda *x: subprocess.run(["git", "-C", str(mono), "-c", "user.email=a@b", "-c", "user.name=a", *x], capture_output=True)
        g("init", "-q"); g("add", "-A"); g("commit", "-qm", "base")
        base = git(mono, "rev-parse", "HEAD")
        (svc / "src/main/java/com/x/ctl/NewController.java").write_text(
            'package com.x.ctl;\nimport com.x.util.Helper;\nimport com.x.util.NotExist;\nimport com.ext.dto.*;\n'
            'import cn.tool.core.convert.Convert;\nimport org.springframework.web.bind.annotation.*;\n'
            '@RestController\n@RequestMapping("svcapi/new")\npublic class NewController { @PostMapping("/_search") public void s(@RequestBody @Valid SendReq req){} }\n', encoding="utf-8")
        (svc / "src/main/java/com/x/ctl/SendReq.java").write_text(
            'package com.x.ctl;\npublic class SendReq {\n private String mobile;\n private String code;\n}\n', encoding="utf-8")
        (svc / "src/main/java/com/x/ctl/BApi.java").write_text(
            'package com.x.ctl;\n@FeignClient(value = "svc-b", path = "/inner/b")\npublic interface BApi {\n'
            ' @PostMapping("/send") void send();\n @PostMapping("/ok") void ok();\n}\n', encoding="utf-8")
        ref = T / "ref"; ref.mkdir()
        (ref / "BController.java").write_text('package com.y;\n@RestController\n@RequestMapping("/inner/b")\n'
                                              'public class BController { @PostMapping("/ok") public void ok(){} }\n', encoding="utf-8")
        (svc / "src/main/java/com/x/ctl/NewEntity.java").write_text(
            'package com.x.ctl;\nimport com.baomidou.mybatisplus.annotation.*;\n@TableName("tb_new")\npublic class NewEntity {\n'
            ' private static final long serialVersionUID = 1L;\n private Long id;\n private String applyNo;\n private String wrongField;\n'
            ' @TableField(exist = false)\n private String ghost;\n}\n', encoding="utf-8")
        (web / "src/api/new.ts").write_text(
            'import { defHttp } from "@/libs/axios";\nimport x from "./missing";\nconst PREFIX = "/svc/svcapi/new";\n'
            'const Api = { A: `${PREFIX}/_search`, B: `${PREFIX}/_nope` };\n', encoding="utf-8")
        (other / "a.txt").write_text("changed", encoding="utf-8")
        (web / "src/api/mode.ts").write_text("// 项目模式映射需内网联调核对\nexport const M = '1';\n", encoding="utf-8")
        g("add", "-A"); g("commit", "-qm", "feat")
        out = T / "out"; (out / "patches").mkdir(parents=True); (out / "S2-02-sql").mkdir(); (out / "S2-03-config").mkdir()
        for f in ["README.md", "S2-07-验证清单.md"]: (out / f).write_text("x", encoding="utf-8")
        (out / "patches/svc-a.patch").write_text("x", encoding="utf-8")
        (out / "patches/svc-z.patch").write_text("x", encoding="utf-8")
        (out / "S2-04-verify").mkdir(); (out / "S2-04-verify/out.txt").write_text("PASS=1", encoding="utf-8")
        (out / "S2-04-verify/cases.json").write_text('[{"ac": "AC-01"}, {"ac": "AC-02"}, {"ac": "AC-06"}]', encoding="utf-8")
        (out / "S2-05-api").mkdir()
        (out / "S2-05-api/接口文档.openapi.json").write_text(json.dumps({"openapi": "3.0.0", "paths": {"/svc/svcapi/new/_search": {}, "/svcapi/ghost": {}}}), encoding="utf-8")
        (out / "S2-05-api/冒烟.postman_collection.json").write_text(json.dumps({"item": [
            {"name": "登录", "request": {"url": "{{baseUrl}}/auth/login"}},
            {"name": "列表", "item": [{"name": "AC-01 列表", "request": {"url": {"raw": "{{baseUrl}}/svc/svcapi/new/_search?x=1"}}},
                                    {"name": "AC-01 反例", "request": {"url": "{{baseUrl}}/svc/svcapi/nope"}}]}]}), encoding="utf-8")
        (out / "S2-02-sql/new.sql").write_text("CREATE TABLE IF NOT EXISTS `tb_new` (\n  `id` bigint NOT NULL,\n  `apply_no` varchar(32)\n) ENGINE=InnoDB;\n", encoding="utf-8")
        data_sql = "INSERT INTO tb_dict (id, code) VALUES (1, 'A');\nDELETE FROM tb_menu WHERE id = 9;\nINSERT INTO tb_menu (id) VALUES (9);\n"
        (out / "S2-02-sql/data.sql").write_text(data_sql, encoding="utf-8")
        big = "".join(f"INSERT INTO tb_cfg (k) SELECT 'k{i}' FROM DUAL WHERE NOT EXISTS (SELECT 1 FROM tb_cfg WHERE k = 'k{i}');\n" for i in range(8))
        (out / "S2-02-sql/big.sql").write_text(big, encoding="utf-8")
        big_hash = hashlib.sha256((out / "S2-02-sql/big.sql").read_bytes()).hexdigest()
        (out / "S2-06-上线手册.md").write_text("```sql\nCREATE TABLE IF NOT EXISTS `tb_new` (\n  `id` bigint NOT NULL,\n  `apply_no` varchar(64)\n) ENGINE=InnoDB;\n```\n"
                                       f"```sql\n{data_sql}```\n执行 S2-02-sql/big.sql sha256:{big_hash}\n"
                                       "```\n  - /svcapi/new/_search\n```\n", encoding="utf-8")
        (out / "S2-03-config/a.diff").write_text("+  - /svcapi/new/_search\n+  - /svcapi/ghost/_x\n", encoding="utf-8")
        s1 = T / "S1"; s1.mkdir()
        (s1 / "S1-03-代码走读.md").write_text("### 接触点\n| 层 | 参照 | 本需求做法（仿 / 改 / 新建） |\n|---|---|---|\n| 入口 | svc-a NewController | 新建 |\n"
                                     "| 其他服务 | svc-b `a.txt` | 改：加接口 |\n| 配置 | web-a-cfg | 不改 |\n", encoding="utf-8")
        (s1 / "S1-07-设计.md").write_text("## 评审摘要\n一屏\nAC-01：a\nAC-02：b\nAC-03：c\nAC-04：d\nAC-05：e\nAC-06：f\nAC-07：g\n\n## 改动面\n| 仓 | 文件 | 改什么 | 单测 |\n|---|---|---|---|\n"
                                   "| svc-a | NewController | x | |\n| web-a | new.ts | x | |\n| web-a | missing.ts | x | |\n| web-b | 路由 | 加路由 | |\n"
                                   "| svc-c | | 不改：无需 | |\n"
                                   "| svc-d | OrderMapper | 修改；precheck 按无改动登记，patch 另验 | |\n"
                                   "\n## 待定决策\n| Q | 问题 | 建议值 | 按建议会怎样 | 状态 |\n|---|---|---|---|---|\n| Q-01 | q | v | 会这样 | 已定 |\n"
                                   "\n### 交付选项\n| 选项 | 建议值 | 依据 | 用户选择 |\n|---|---|---|---|\n| 环境形态 | CI | 手册 | CI |\n| 交付范围 | 仅后端 | 仓 | 仅后端 |\n", encoding="utf-8")
        (out / "S2-08-审查.md").write_text("## AC 实现映射\n| AC | 实现位置 | 结论 |\n|---|---|---|\n"
                                    "| AC-01 | svc-a/src/main/java/com/x/ctl/NewController.java:3 | 已实现 |\n"
                                    "| AC-02 | src/main/java/com/x/ctl/Nope.java:9 | 已实现 |\n"
                                    "| AC-04 | 由平台承载 | 平台承载 |\n| AC-05 | 环境侧 | 环境登记 |\n| AC-06 | 页面 | 本期不做（D-01） |\n| AC-07 | 页面 | 本期不做（D-01） |\n\n## 问题\n| 编号 | 类别 | 问题 | 位置 | 状态 |\n|---|---|---|---|---|\n"
                                    "| P-01 | 需求不符 | 审批可绕过 | | 待修 |\n| P-02 | 安全 | 身份取自请求 | | 后续优化 |\n", encoding="utf-8")
        cfgp = T / "precheck.json"; cfgp.write_text(json.dumps({"api_prefixes": ["/svcapi/"], "table_prefixes": ["tb_"],
                                                                "ref_dirs": [str(ref)]}), encoding="utf-8")
        ns = lambda **kw: argparse.Namespace(**{"out": str(out), "handoff": None, "repo": [str(svc), str(web)],
                                               "base": [f"svc-a={base}", f"web-a={base}"], "no_patch": False, "allow_dirty": False,
                                               "write_manifest": True, "config": str(cfgp), **kw})
        rep = run(ns()); text, ok = rep.finish()
        msgs = [(lv, c, m) for lv, c, m in rep.issues]
        has = lambda lv, c, s="": any(l == lv and k == c and s in m for l, k, m in msgs)
        failures = []
        for want, s in (("J1", "NotExist"), ("J2", "wrong_field"), ("F1", "missing"), ("Q1", ""), ("Q2", "ghost")):
            if not has("ERROR", want, s): failures.append(f"期望 ERROR {want} 未触发")
        if not has("WARN", "N1", "ghost"): failures.append("期望 WARN N1 未触发")
        if not has("WARN", "F2", "_nope"): failures.append("期望 WARN F2 未触发")
        if not has("WARN", "U1", "svc-b"): failures.append("期望 WARN U1（未检查目录 svc-b）未触发")
        if not has("WARN", "J1", "com.ext.dto.*"): failures.append("外部通配导入且无同组依赖时应 WARN")
        if has("WARN", "J1", "Convert"): failures.append("同组依赖已在用（cn.tool.core）仍报 WARN")
        if any(c == "J2" and ("ghost" in m or "serialVersionUID" in m) for _, c, m in msgs): failures.append("exist=false / static 字段未被排除")
        if any(c == "F2" and m.endswith("/svc/svcapi/new 在 Controller 映射中未找到") for _, c, m in msgs): failures.append("前缀常量被当成请求路径")
        if any(c == "Q2" and "_search" in m for _, c, m in msgs): failures.append("手册已含的配置行被误报")
        for want, s in (("Q3", "tb_dict"), ("T1", "mode.ts"), ("A1", "/svcapi/ghost"), ("A1", "/svc/svcapi/nope"), ("A1", "AC-02"), ("C1", "web-b"), ("C1", "svc-d"), ("C1", "missing.ts"), ("R1", "AC-03"), ("R1", "Nope.java"), ("R1", "P-01"), ("R1", "AC-04"),
                        ("R1", "AC-05"), ("R1", "AC-06"), ("C2", "svc-b"), ("U2", "svc-z"), ("H2", "out.txt"), ("V1", "AC-01"), ("X1", "正文.md")):
            if not has("ERROR", want, s): failures.append(f"期望 ERROR {want}（{s}）未触发")
        for c, s in (("Q3", "tb_menu"), ("Q1", "big.sql"), ("C1", "svc-a"), ("C1", "svc-c"), ("C1", "new.ts"), ("R1", "AC-01"), ("R1", "AC-07"), ("R1", "P-02"), ("N2", "/inner/b/ok"),
                     ("C2", "svc-a"), ("C2", "web-a"), ("U2", "svc-a"), ("D1", ""), ("A1", "_search"), ("A1", "login"), ("A1", "AC-01,"), ("V1", "AC-07")):
            if has("ERROR", c, s) or has("WARN", c, s): failures.append(f"{c}（{s}）不应报")
        if not has("WARN", "N2", "/inner/b/send"): failures.append("期望 WARN N2 未触发")
        if not has("WARN", "S1", "mobile"): failures.append("期望 WARN S1 未触发")
        if rep.counts.get("svc-a", {}).get("java") != 4 or rep.counts.get("web-a", {}).get("fe") != 2: failures.append(f"文件计数错误 {rep.counts}")
        if ok: failures.append("有 ERROR 却 PASS")
        rep2 = run(ns(repo=[str(mono)], base=[f"mono={git(mono, 'rev-parse', 'HEAD')}"])); _, ok2 = rep2.finish()
        if ok2 or not any(c == "E0" for _, c, _ in rep2.issues): failures.append("无改动时应 ERROR E0")
        rep3 = run(ns(no_patch=True, write_manifest=False))
        if any(c == "M1" for _, c, _ in rep3.issues): failures.append("--no-patch 不应检查 manifest")
        s1b = T / "req/S1"; s1b.mkdir(parents=True); (T / "_project").mkdir()
        (s1b / "S1-06-交互流程.md").write_text("| /x/b | 不画 |\nA->>B: GET /x/a\n", encoding="utf-8")
        (s1b / "S1-04-设计说明.md").write_text("## 接口契约\n| 方法 | 路径 | 调用者 |\n|---|---|---|\n| GET | /x/a | u |\n| POST | `/x/b?k=` | u |\n\n## 改动面\n| 仓 | 文件 |\n|---|---|\n| svc-a | x |\n", encoding="utf-8")
        (s1b / "S1-05-待定决策.md").write_text("### 交付选项\n| 选项 | 建议值 | 用户选择 |\n|---|---|---|\n| 环境形态 | CI | CI |\n"
                                        "| 交付范围 | 仅后端 | 仅后端 |\n| 冒烟集合 | 产出 | （待用户回复） |\n", encoding="utf-8")
        (T / "_project/S0-01-项目手册.md").write_text("## 1. 仓库地图\n", encoding="utf-8")
        rep4 = run(ns(s1=str(s1b), handoff=str(T / "ho/new")))
        d1 = [m for _, c, m in rep4.issues if c == "D1"]
        for s in ("评审摘要", "调用者/数据范围/状态前置", "按建议会怎样", "术语映射", "冒烟集合", "/x/b"):
            if not any(s in m for m in d1): failures.append(f"期望 D1（{s}）未触发")
        if any("/x/a" in m or "「环境形态」" in m or "「交付范围」" in m for m in d1): failures.append("已画的接口或已回复的选项被误报")
        if not (T / "ho/new/manifest.txt").exists(): failures.append("handoff 目录不存在时应自动创建再写 manifest")
        sop = T / "sop"; (sop / "a").mkdir(parents=True); (sop / "b").mkdir()
        (sop / "a/SKILL.md").write_text("x\n" * 125 + "用 acme-pay 发起，" + "20" + "26-01-15 修订\n", encoding="utf-8")
        (sop / "b/x.md").write_text("通用流程\n", encoding="utf-8")
        ltext, lok = lint_sop(sop, {**DEFAULT_CONFIG, "project_terms": ["ACME-pay"]})
        for c in ("L1", "L2", "L3"):
            if f"[{c}]" not in ltext: failures.append(f"期望 lint {c} 未触发")
        if lok or "b/x.md" in ltext: failures.append("lint 结果错误")
        ks = T / "k1s1"; ks.mkdir()
        (ks / "S1-07-设计.md").write_text("类型：小改动\n", encoding="utf-8")
        (ks / "S1-03-代码走读.md").write_text("## 平台登记数据取样\n| 登记项 | 本需求要加/依赖的行 |\n|---|---|\n| 字典 | 新增一行 |\n", encoding="utf-8")
        kr = Report(); check_type(ks, [out / "S2-02-sql/new.sql"], 0, kr)
        if not any(c == "K1" and "新建表" in m for _, c, m in kr.issues): failures.append("K1 新建表未触发")
        if not any(c == "K1" and "平台登记" in m for _, c, m in kr.issues): failures.append("K1 平台登记未触发")
        (ks / "S1-07-设计.md").write_text("类型：纯配置\n", encoding="utf-8")
        kr3 = Report(); check_type(ks, [], 2, kr3)
        if not any(c == "K1" and "代码改动" in m for _, c, m in kr3.issues): failures.append("K1 纯配置有代码未触发")
        ds = T / "d2"; ds.mkdir()
        (ds / "S1-01-决定表.md").write_text("| 编号 | 原句 | 来源 | 状态 |\n|---|---|---|---|\n| D-01 | 要做 | 原文 | 有效 |\n| D-02 | 不做 | 原文 | 作废 |\n", encoding="utf-8")
        (ds / "S1-02-需求与验收.md").write_text("| REQ | 依据 |\n|---|---|\n| REQ-01 | D-09 |\n", encoding="utf-8")
        dr = Report(); check_decisions(ds, dr)
        if not any(c == "D2" and "D-01" in m for _, c, m in dr.issues): failures.append("D2 有效项未触发")
        if any(c == "D2" and "D-02" in m for _, c, m in dr.issues): failures.append("作废的 D-02 不应要求验收")
        (ds / "S1-04-设计说明.md").write_text("x\n", encoding="utf-8"); (ds / "S1-01-决定表.md").unlink()
        drb = Report(); check_decisions(ds, drb)
        if not any(c == "D2" and "S1-01-决定表.md" in m for _, c, m in drb.issues): failures.append("D2 缺决定表未触发")
        ing = T / "ing"
        docx = T / "sample.docx"
        png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 16
        xlsx_buf = io.BytesIO()
        import openpyxl
        wb = openpyxl.Workbook(); wb.active.title = "分类"; wb.active.append(["名称", "说明"]); wb.active.append(["甲", "要验证"])
        wb.save(xlsx_buf)
        wml = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        doc_xml = (
            f'<?xml version="1.0" encoding="UTF-8"?>'
            f'<w:document xmlns:w="{wml}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:body>'
            f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>查看前要验证</w:t></w:r></w:p>'
            f'<w:p><w:r><w:t></w:t></w:r>'
            f'<w:r><w:drawing><a:blip r:embed="rIdImg"/></w:drawing></w:r></w:p>'
            f'<w:tbl><w:tr><w:tc><w:p><w:r><w:t>列甲</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
            f'</w:body></w:document>')
        with zipfile.ZipFile(docx, "w") as z:
            z.writestr("word/document.xml", doc_xml.encode("utf-8"))
            z.writestr("word/_rels/document.xml.rels",
                       '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                       '<Relationship Id="rIdImg" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/>'
                       '<Relationship Id="rIdHdr" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" Target="header1.xml"/>'
                       '</Relationships>')
            z.writestr("word/media/image1.png", png)
            z.writestr("word/header1.xml",
                       f'<?xml version="1.0" encoding="UTF-8"?><w:hdr xmlns:w="{wml}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                       f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:p><w:r><w:drawing><a:blip r:embed="rIdH"/></w:drawing></w:r></w:p></w:hdr>')
            z.writestr("word/_rels/header1.xml.rels",
                       '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                       '<Relationship Id="rIdH" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/logo.png"/>'
                       '</Relationships>')
            z.writestr("word/media/logo.png", png)
            z.writestr("word/embeddings/a.xlsx", xlsx_buf.getvalue())
        gbk = T / "note.txt"; gbk.write_bytes("仅预览".encode("gbk"))
        held = []
        def _cap(*args, **kwargs):
            held.append(" ".join(str(x) for x in args))
        import builtins
        real_print = builtins.print
        try:
            builtins.print = _cap
            code = ingest([str(docx), str(gbk)], ing)
        except Exception as e:
            failures.append(f"ingest 异常 {e}")
            code = 1
        finally:
            builtins.print = real_print
        line = "\n".join(held)
        if code != 0 or line.strip() != "RESULT: PASS": failures.append(f"ingest 结果异常 {code} {line!r}")
        if any(ord(ch) > 127 for ch in line): failures.append("ingest 标准输出含非 ASCII")
        body = (ing / "正文.md").read_text(encoding="utf-8") if (ing / "正文.md").exists() else ""
        if "查看前要验证" not in body or "仅预览" not in body or "列甲" not in body or "要验证" not in body:
            failures.append("ingest 正文丢了中文、表格或内嵌表")
        if "〔图 media/b1.png〕" not in body or not (ing / "media/b1.png").exists(): failures.append("ingest 没有抽出图片")
        man0 = (ing / "清单.md").read_text(encoding="utf-8")
        if "| media/b2.png | 页眉 | 忽略 | 页眉 |" not in man0: failures.append("页眉图应标忽略")
        s1ing = T / "s1ing"; (s1ing / "S1-00-原文").mkdir(parents=True)
        for f in ing.iterdir():
            target = s1ing / "S1-00-原文" / f.name
            if f.is_dir(): shutil.copytree(f, target, dirs_exist_ok=True)
            else: shutil.copy(f, target)
        src = Report(); check_source(s1ing, src)
        if not any(c == "X1" and "待看" in m for _, c, m in src.issues): failures.append("待看的图应被 X1 拦住")
        man = (s1ing / "S1-00-原文/清单.md").read_text(encoding="utf-8").replace("| media/b1.png | 正文 | 待看 |  |", "| media/b1.png | 正文 | 已读 |  |")
        (s1ing / "S1-00-原文/清单.md").write_text(man, encoding="utf-8")
        src2 = Report(); check_source(s1ing, src2)
        if any(c == "X1" for _, c, _ in src2.issues): failures.append("已读且 UTF-8 的原文不应报 X1")
        empty = T / "emptyS1"; (empty / "S1-00-原文").mkdir(parents=True)
        (empty / "S1-00-原文/正文.md").write_text("## 原型\n\n〔图 media/b1.png〕\n", encoding="utf-8")
        (empty / "S1-00-原文/media").mkdir()
        (empty / "S1-00-原文/media/b1.png").write_bytes(png)
        (empty / "S1-00-原文/清单.md").write_text("## 图片\n| 文件 | 位置 | 状态 | 原因 |\n|---|---|---|---|\n| media/b1.png | 原型 | 忽略 | 装饰 |\n", encoding="utf-8")
        er = Report(); check_source(empty, er)
        if not any(c == "X1" and "不能忽略" in m for _, c, m in er.issues): failures.append("空节旁的图被忽略时应报 X1")
        emf = T / "vec.docx"
        with zipfile.ZipFile(emf, "w") as z:
            z.writestr("word/document.xml",
                       f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{wml}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                       f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:body><w:p><w:r><w:t>有字</w:t></w:r>'
                       f'<w:r><w:drawing><a:blip r:embed="rIdE"/></w:drawing></w:r></w:p></w:body></w:document>')
            z.writestr("word/_rels/document.xml.rels",
                       '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                       '<Relationship Id="rIdE" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/a.emf"/>'
                       '</Relationships>')
            z.writestr("word/media/a.emf", b"EMF")
        builtins.print = lambda *a, **k: None
        try:
            ingest([str(emf)], T / "vemf")
        finally:
            builtins.print = real_print
        vman = (T / "vemf/清单.md").read_text(encoding="utf-8") if (T / "vemf/清单.md").exists() else ""
        if "| media/b1.emf | 正文 | 无法看 |" not in vman: failures.append("emf 应标无法看且保留扩展名")
        blank = T / "blank.pdf"
        blank.write_bytes(b"%PDF-1.1\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n2 0 obj<</Type/Pages/Count 1/Kids[3 0 R]>>endobj\n3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 3 3]>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF\n")
        bout = T / "blankout"
        builtins.print = lambda *a, **k: None
        try:
            bcode = ingest([str(blank)], bout)
        finally:
            builtins.print = real_print
        if bcode == 0 or (bout / "正文.md").exists(): failures.append("整份无文本层不应写出正文.md")
        bman = (bout / "清单.md").read_text(encoding="utf-8") if (bout / "清单.md").exists() else ""
        if "整份无文本层" not in bman: failures.append("整份无文本层应记入清单，不能当成已落盘原文")
        print(text); print("---")
        if failures:
            print("SELFTEST FAIL"); [print(" -", f) for f in failures]; return 1
        print("SELFTEST PASS（J1 J2 F1 F2 Q1 Q2 Q3 N1 N2 S1 C1 C2 D1 R1 H2 T1 A1 U1 U2 E0 M1 L1-L3 均触发；防护项均不误报）"); return 0
    finally:
        shutil.rmtree(T, ignore_errors=True)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", help="文档目录（S2/：S2-06-上线手册、S2-07-验证清单、S2-02-sql/、S2-03-config/）")
    ap.add_argument("--handoff", help="改动包目录（patches/、README、manifest）；不给且非 --no-patch 时等于 --out")
    ap.add_argument("--repo", action="append", default=[])
    ap.add_argument("--base", action="append", default=[], help="目录名=基线commit")
    ap.add_argument("--no-patch", action="store_true", help="分支/PR 形态：不要求 patches/README/manifest")
    ap.add_argument("--allow-dirty", action="store_true")
    ap.add_argument("--write-manifest", action="store_true")
    ap.add_argument("--config", help="项目参数 JSON（建议 artifacts/_project/S0-02-precheck.json）")
    ap.add_argument("--s1", help="S1 目录；默认 --out 的同级 S1/")
    ap.add_argument("--lint-sop", help="检查 sop-v2 自身：篇幅、日期、project_terms（见 SOP约束.md）")
    ap.add_argument("--ingest", nargs="+", help="原文文件。与 --out（S1/S1-00-原文目录）合用，落成 UTF-8 正文")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"): sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if a.selftest: sys.exit(selftest())
    if a.lint_sop:
        text, ok = lint_sop(a.lint_sop, load_cfg(a.config)); print(text); sys.exit(0 if ok else 1)
    if a.ingest:
        if not a.out: ap.error("--ingest 需要 --out")
        sys.exit(ingest(a.ingest, a.out))
    if not a.out or not a.repo: ap.error("--out 与至少一个 --repo 必填")
    rep = run(a); text, ok = rep.finish()
    print(text)
    rdir = Path(a.handoff) if a.handoff else Path(a.out)
    if rdir.exists(): (rdir / "precheck-report.txt").write_text(text + "\n", encoding="utf-8")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
